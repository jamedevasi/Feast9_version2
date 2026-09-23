from datetime import date, timedelta

import pytest

from app import db
from tests.conftest import get_csrf
from tests.test_cases import _create_case
from tests.test_roles import _create_user, _login, _logout

TODAY = date.today().isoformat()
LONG_AGO = "2000-01-01"


def _case_for(client, patient_id, **overrides):
    resp, doctor_id = _create_case(client, patient_id, **overrides)
    case_url = resp.headers["Location"]
    case_id = int(case_url.rstrip("/").rsplit("/", 1)[-1])
    return case_id, case_url, doctor_id


def _add_payment(client, case_id, case_url, amount, payment_date=None):
    token = get_csrf(client, case_url)
    client.post(
        f"/cases/{case_id}/payments",
        data={
            "payment_date": payment_date or TODAY,
            "amount": str(amount),
            "method": "Cash",
            "csrf_token": token,
        },
    )


def _add_visit_note(client, case_id, case_url, visit_date):
    token = get_csrf(client, case_url)
    client.post(
        f"/cases/{case_id}/visit-notes",
        data={"note": "Checkup", "visit_date": visit_date, "csrf_token": token},
    )


def test_new_cases_and_patients_counts(logged_in_client, patient_id):
    _case_for(logged_in_client, patient_id, title="Reports Case")
    assert db.get_new_cases_count(TODAY, TODAY) >= 1
    assert db.get_new_cases_count(LONG_AGO, "2000-01-02") == 0
    assert db.get_new_patients_count(TODAY, TODAY) >= 1
    assert db.get_new_patients_count(LONG_AGO, "2000-01-02") == 0


def test_cases_closed_count_and_list(logged_in_client, patient_id):
    case_id, case_url, _ = _case_for(logged_in_client, patient_id, title="Close Me")
    token = get_csrf(logged_in_client, case_url)
    logged_in_client.post(f"/cases/{case_id}/close", data={"csrf_token": token})

    assert db.get_cases_closed_count(TODAY, TODAY) >= 1
    closed = db.get_cases_closed_in_range(TODAY, TODAY)
    assert any(c["id"] == case_id for c in closed)
    assert db.get_cases_closed_count(LONG_AGO, "2000-01-02") == 0


def test_revenue_collected_and_overview(logged_in_client, patient_id):
    case_id, case_url, _ = _case_for(logged_in_client, patient_id, total_cost="3000")
    _add_payment(logged_in_client, case_id, case_url, 1200)

    assert db.get_revenue_collected(TODAY, TODAY) >= 1200
    overview = db.get_revenue_overview()
    assert overview["billed"] >= 3000
    assert overview["collected"] >= 1200
    assert overview["outstanding"] == pytest.approx(overview["billed"] - overview["collected"])


def test_pending_payments_lists_only_unpaid_balance(logged_in_client, patient_id):
    case_id, case_url, _ = _case_for(logged_in_client, patient_id, title="Pending Case", total_cost="2000")
    _add_payment(logged_in_client, case_id, case_url, 500)

    pending = db.get_pending_payments()
    row = next(r for r in pending if r["case_id"] == case_id)
    assert row["balance"] == 1500

    case_id2, case_url2, _ = _case_for(logged_in_client, patient_id, title="Paid Case", total_cost="800")
    _add_payment(logged_in_client, case_id2, case_url2, 800)
    pending2 = db.get_pending_payments()
    assert not any(r["case_id"] == case_id2 for r in pending2)


def test_doctor_revenue_and_collection_rate(logged_in_client, patient_id):
    case_id, case_url, doctor_id = _case_for(logged_in_client, patient_id, total_cost="1000")
    _add_payment(logged_in_client, case_id, case_url, 500)

    revenue = db.get_doctor_revenue_by_period(TODAY, TODAY)
    row = next(r for r in revenue if r["doctor_id"] == doctor_id)
    assert row["billed"] >= 1000
    assert row["collected"] >= 500
    assert row["collection_rate"] == pytest.approx(row["collected"] / row["billed"] * 100)


def test_patient_retention_lapsed_and_all_clear(logged_in_client, patient_id):
    retention = db.get_patient_retention()
    assert retention["total_patients"] >= 1
    assert not any(p["id"] == patient_id for p in retention["lapsed"])

    case_id, case_url, _ = _case_for(logged_in_client, patient_id)
    old_date = (date.today() - timedelta(days=250)).isoformat()
    _add_visit_note(logged_in_client, case_id, case_url, old_date)

    retention2 = db.get_patient_retention(months=6)
    assert any(p["id"] == patient_id for p in retention2["lapsed"])


def test_reports_empty_state_all_clear(logged_in_client):
    resp = logged_in_client.get("/reports/")
    assert resp.status_code == 200
    assert b"All clear" in resp.data


def test_reports_page_renders_all_sections(logged_in_client, patient_id):
    case_id, case_url, _ = _case_for(logged_in_client, patient_id, title="Full Report Case", total_cost="1000")
    _add_payment(logged_in_client, case_id, case_url, 400)
    resp = logged_in_client.get("/reports/")
    body = resp.data.decode()
    assert "Revenue Overview" in body
    assert "Cases Ledger" in body  # merged Pending/Payments/Closed table (default: pending)
    assert "Doctor-wise Revenue" in body
    assert "Patient Retention" in body
    assert "Full Report Case" in body


def test_reports_filter_form_has_explicit_action(logged_in_client):
    resp = logged_in_client.get("/reports/")
    assert b'action="/reports/"' in resp.data


def test_reports_pdf_download(logged_in_client, patient_id):
    resp = logged_in_client.get("/reports/report.pdf")
    assert resp.status_code == 200
    assert resp.headers["Content-Type"] == "application/pdf"
    assert resp.data[:4] == b"%PDF"


def test_download_ledger_excel_matches_current_view(logged_in_client, patient_id):
    from io import BytesIO
    from openpyxl import load_workbook

    case_id, case_url, _ = _case_for(logged_in_client, patient_id, title="Excel Case", total_cost="1000")
    _add_payment(logged_in_client, case_id, case_url, 250)
    resp = logged_in_client.get("/reports/ledger.xlsx?view=payments&group=method")
    assert resp.status_code == 200
    assert resp.headers["Content-Type"].startswith("application/vnd.openxmlformats")
    ws = load_workbook(BytesIO(resp.data)).active
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    assert rows[0][:5] == ["Date", "Patient", "Case", "Doctor", "Amount"]
    assert any(r[2] == "Excel Case" and r[4] == 250 for r in rows)
    assert rows[-1][0].startswith("Total") and rows[-1][4] == 250
    assert any(str(r[0]).startswith("Subtotal") for r in rows)


def test_receptionist_blocked_from_reports(logged_in_client, patient_id):
    _create_user(logged_in_client, "repreports", "receptionist")
    _logout(logged_in_client)
    _login(logged_in_client, "repreports")

    assert logged_in_client.get("/reports/").status_code == 403
    assert logged_in_client.get("/reports/report.pdf").status_code == 403
    assert logged_in_client.get("/reports/ledger.xlsx").status_code == 403
    assert logged_in_client.get("/reports/ledger.xlsx?view=payments").status_code == 403

    dash = logged_in_client.get("/dashboard")
    assert b'href="/reports/"' not in dash.data


def test_reports_shows_profitability_summary_not_full_table(logged_in_client, patient_id):
    """User instruction (2026-09-18): the Financial Assessment expense table must stay on its
    own page — Reports gets a compact all-time summary + link only, never the editable table,
    so the period-filtered report view doesn't get cluttered with a second, all-time table."""
    case_id, case_url, _ = _case_for(logged_in_client, patient_id, title="Profitability Case", total_cost="1000")
    _add_payment(logged_in_client, case_id, case_url, 1000)

    resp = logged_in_client.get("/reports/")
    assert b"Profitability" in resp.data
    assert b"Financial Assessment" in resp.data
    # The editable per-case expense inputs (Financial Assessment's own table) must never
    # render inside Reports — only a rolled-up summary and a link to the dedicated page.
    assert b'name="lab_amount_' not in resp.data
    assert b'name="consultant_fee_' not in resp.data
    assert b'name="misc_expense_' not in resp.data


def test_ledger_default_is_pending_payments(logged_in_client, patient_id):
    _case_for(logged_in_client, patient_id, title="Owes Money", total_cost="1000")
    paid_id, paid_url, _ = _case_for(logged_in_client, patient_id, title="Settled Up", total_cost="300")
    _add_payment(logged_in_client, paid_id, paid_url, 300)
    body = logged_in_client.get("/reports/").data.decode()
    assert "Owes Money" in body
    assert "Settled Up" not in body
    # "Any balance" (present-but-empty) is a real choice, not a fallback to the default.
    body = logged_in_client.get("/reports/?view=cases&balance=").data.decode()
    assert "Settled Up" in body


def test_ledger_cases_closed_in_period(logged_in_client, patient_id):
    case_id, case_url, _ = _case_for(logged_in_client, patient_id, title="Closed One", total_cost="0")
    token = get_csrf(logged_in_client, case_url)
    logged_in_client.post(f"/cases/{case_id}/close", data={"csrf_token": token})
    _case_for(logged_in_client, patient_id, title="Still Open", total_cost="0")
    body = logged_in_client.get(
        f"/reports/?view=cases&balance=&period_on=closed&start={TODAY}&end={TODAY}"
    ).data.decode()
    assert "Closed One" in body
    assert "Still Open" not in body


def test_ledger_payments_view_grouped_with_subtotals(logged_in_client, patient_id):
    case_id, case_url, _ = _case_for(logged_in_client, patient_id, title="Grouped Case", total_cost="5000")
    _add_payment(logged_in_client, case_id, case_url, 100)
    _add_payment(logged_in_client, case_id, case_url, 250)
    body = logged_in_client.get("/reports/?view=payments&group=patient").data.decode()
    assert "Payments Ledger" in body
    assert 'class="ledger-group"' in body
    assert "₹350.00" in body  # the patient's subtotal (and the grand total)


def test_ledger_ignores_unknown_options(logged_in_client, patient_id):
    resp = logged_in_client.get("/reports/?view=bogus&group=drop_table&balance=x&period_on=y")
    assert resp.status_code == 200
    assert b"Cases Ledger" in resp.data
