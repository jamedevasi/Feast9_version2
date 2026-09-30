"""Security review fixes (2026-09-30): a generated session-signing key instead of a publicly
known one, no live formulas in any Excel export, and no browser caching of app pages."""
import io

from openpyxl import load_workbook

from app import config as app_config
from app import create_app, db

EVIL = '=HYPERLINK("http://evil.example","click")'


# ── 1. Session signing key ─────────────────────────────────────────────────

def test_generated_secret_key_is_random_and_persistent(app, monkeypatch):
    monkeypatch.setattr(app_config, "SECRET_KEY", "")
    key = create_app().config["SECRET_KEY"]
    assert len(key) == 64 and key not in app_config.KNOWN_PUBLIC_SECRET_KEYS
    with open(app_config.secret_key_file(), encoding="utf-8") as f:
        assert f.read().strip() == key
    assert create_app().config["SECRET_KEY"] == key  # same key on the next start — logins survive a restart


def test_publicly_known_secret_keys_are_never_used(app, monkeypatch):
    for public in app_config.KNOWN_PUBLIC_SECRET_KEYS:
        monkeypatch.setattr(app_config, "SECRET_KEY", public)
        assert create_app().config["SECRET_KEY"] not in app_config.KNOWN_PUBLIC_SECRET_KEYS


def test_explicit_strong_secret_key_wins(app, monkeypatch):
    strong = "c" * 64
    monkeypatch.setattr(app_config, "SECRET_KEY", strong)
    assert create_app().config["SECRET_KEY"] == strong


def test_session_cookie_signed_with_public_key_is_rejected(logged_in_client):
    """A forged 'admin' cookie made with the old published key must not log anyone in."""
    from flask.sessions import SecureCookieSessionInterface
    from flask import Flask
    forger = Flask("forger")
    forger.secret_key = "local-dev-key"
    cookie = SecureCookieSessionInterface().get_signing_serializer(forger).dumps(
        {"admin_id": 1, "role": "admin", "reauth_at": "2099-01-01T00:00:00+00:00"}
    )
    logged_in_client.delete_cookie("session")  # drop the real session
    logged_in_client.set_cookie("session", cookie)
    resp = logged_in_client.get("/backup/")
    assert resp.status_code == 302 and "/login" in resp.headers["Location"]


# ── 2. No live formulas in Excel exports ───────────────────────────────────

def _case_for(name):
    doctor_id = db.add_doctor("Dr Test")
    pid = db.add_patient({"name": name, "sex": "Female", "dpdp_notice_accepted": 1})
    return db.add_case({"patient_id": pid, "title": EVIL, "doctor_id": doctor_id, "total_cost": 500})


def _cells(xlsx_bytes):
    ws = load_workbook(io.BytesIO(xlsx_bytes)).active
    return [c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]


def test_reports_ledger_export_never_contains_formulas(logged_in_client):
    _case_for(EVIL)
    resp = logged_in_client.get("/reports/ledger.xlsx?group=patient")
    cells = _cells(resp.data)
    assert cells, "the test data should have reached the export"
    assert all(c.data_type == "s" for c in cells)


def test_financial_assessment_export_never_contains_formulas(logged_in_client):
    _case_for(EVIL)
    cells = _cells(logged_in_client.get("/financial-assessment/export.xlsx").data)
    assert cells and all(c.data_type == "s" for c in cells)


def test_template_doctor_list_never_contains_formulas(logged_in_client):
    db.add_doctor(EVIL)
    wb = load_workbook(io.BytesIO(logged_in_client.get("/import/template.xlsx").data))
    [cell] = [c for c in wb["Lists"]["A"] if c.value == EVIL]
    assert cell.data_type == "s"


# ── 3. No browser caching of app pages ─────────────────────────────────────

def test_sensitive_pages_and_downloads_are_not_cached(logged_in_client, patient_id):
    from app import backup as backup_module
    backup_module.create_key()
    for path in (f"/patients/{patient_id}", "/backup/key", "/backup/", "/import/export.xlsx", "/login"):
        resp = logged_in_client.get(path)
        assert resp.headers.get("Cache-Control") == "no-store", path


def test_public_static_files_stay_cacheable(client):
    resp = client.get("/static/style.css")
    assert resp.status_code == 200
    assert resp.headers.get("Cache-Control") != "no-store"
