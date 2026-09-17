import io

from openpyxl import Workbook, load_workbook

from app import db
from app.excel_import import TEMPLATE_HEADERS
from tests.conftest import get_csrf
from tests.test_roles import _create_user, _login, _logout


def _build_workbook(rows):
    wb = Workbook()
    ws = wb.active
    ws.append(TEMPLATE_HEADERS)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


VALID_ROW = [
    "Ananya Sharma", "1990-04-12", "", "Female", "9876543210", "a@b.com", "1 Main St",
    "Diabetes", "Penicillin", "No", "No", "", "", "", "Yes", "", "No", "", "", "",
]
MINOR_WITH_GUARDIAN = [
    "Minor Kid", "2015-01-01", "", "Male", "", "", "", "", "", "No", "No",
    "Parent Name", "Mother", "9876500001", "Yes", "", "No", "Parent Name", "Mother", "9876500001",
]
MINOR_WITHOUT_GUARDIAN = [
    "Unguarded Kid", "2015-01-01", "", "Male", "", "", "", "", "", "No", "No",
    "", "", "", "Yes", "", "No", "", "", "",
]
NO_CONSENT_ROW = [
    "No Consent Person", "1980-01-01", "", "Male", "", "", "", "", "", "No", "No",
    "", "", "", "No", "", "No", "", "", "",
]
BAD_SEX_ROW = [
    "Bad Sex Person", "1980-01-01", "", "Other", "", "", "", "", "", "No", "No",
    "", "", "", "Yes", "", "No", "", "", "",
]


def _upload(client, content, filename="patients.xlsx"):
    token = get_csrf(client, "/import/")
    return client.post(
        "/import/patients",
        data={"file": (io.BytesIO(content), filename), "csrf_token": token},
        content_type="multipart/form-data",
    )


def test_download_template_has_expected_headers(logged_in_client):
    resp = logged_in_client.get("/import/template.xlsx")
    assert resp.status_code == 200
    assert resp.headers["Content-Type"].startswith("application/vnd.openxmlformats")

    wb = load_workbook(io.BytesIO(resp.data))
    ws = wb.active
    header_row = next(ws.iter_rows(values_only=True))
    assert list(header_row) == TEMPLATE_HEADERS


def test_import_valid_rows(logged_in_client):
    content = _build_workbook([VALID_ROW, MINOR_WITH_GUARDIAN])
    resp = _upload(logged_in_client, content)
    assert resp.status_code == 200
    assert b"2</strong> patient(s) imported" in resp.data

    patients = db.list_patients()
    names = {p["name"] for p in patients}
    assert "Ananya Sharma" in names
    assert "Minor Kid" in names
    for p in patients:
        assert p["is_historic_import"] == 1
        assert p["dpdp_notice_accepted"] == 1


def test_dpdp_not_accepted_row_skipped(logged_in_client):
    content = _build_workbook([NO_CONSENT_ROW])
    resp = _upload(logged_in_client, content)
    body = resp.data.decode()
    assert "0</strong> patient(s) imported" in body
    assert "DPDP Notice Accepted must be Yes" in body
    assert db.list_patients() == []


def test_invalid_sex_row_skipped(logged_in_client):
    content = _build_workbook([BAD_SEX_ROW])
    resp = _upload(logged_in_client, content)
    assert "Sex must be Male or Female" in resp.data.decode()
    assert db.list_patients() == []


def test_minor_without_guardian_skipped(logged_in_client):
    content = _build_workbook([MINOR_WITHOUT_GUARDIAN])
    resp = _upload(logged_in_client, content)
    body = resp.data.decode()
    assert "Guardian name is required" in body
    assert db.list_patients() == []


def test_mixed_valid_and_invalid_rows(logged_in_client):
    content = _build_workbook([VALID_ROW, NO_CONSENT_ROW, BAD_SEX_ROW])
    resp = _upload(logged_in_client, content)
    body = resp.data.decode()
    assert "1</strong> patient(s) imported" in body
    assert "2</strong> row(s) skipped" in body or "2" in body
    assert len(db.list_patients()) == 1


def test_missing_required_columns_rejected(logged_in_client):
    wb = Workbook()
    ws = wb.active
    ws.append(["Name", "Mobile"])  # missing Sex and DPDP Notice Accepted
    ws.append(["Someone", "9876543210"])
    buf = io.BytesIO()
    wb.save(buf)

    resp = _upload(logged_in_client, buf.getvalue())
    assert resp.status_code == 302
    follow = logged_in_client.get("/import/")
    assert b"Missing required column" in follow.data
    assert db.list_patients() == []


def test_non_xlsx_file_rejected(logged_in_client):
    resp = _upload(logged_in_client, b"just some text", filename="patients.txt")
    assert resp.status_code == 302


def test_import_is_audited(logged_in_client):
    content = _build_workbook([VALID_ROW])
    _upload(logged_in_client, content)
    actions = [e["action"] for e in db.list_audit_log()]
    assert "patients_bulk_imported" in actions


def test_imported_patient_shows_historic_badge_on_detail_page(logged_in_client):
    content = _build_workbook([VALID_ROW])
    _upload(logged_in_client, content)
    patient_id = db.list_patients()[0]["id"]
    resp = logged_in_client.get(f"/patients/{patient_id}")
    assert b"Historic Import" in resp.data


def test_import_requires_reauth_when_session_stale(logged_in_client):
    with logged_in_client.session_transaction() as sess:
        sess["reauth_at"] = "2020-01-01T00:00:00"

    content = _build_workbook([VALID_ROW])
    resp = _upload(logged_in_client, content)
    assert resp.status_code == 302
    assert "/reauth" in resp.headers["Location"]
    assert db.list_patients() == []


def test_non_admin_blocked_from_import(logged_in_client):
    _create_user(logged_in_client, "importrecep", "receptionist")
    _logout(logged_in_client)
    _login(logged_in_client, "importrecep")

    assert logged_in_client.get("/import/").status_code == 403
    assert logged_in_client.get("/import/template.xlsx").status_code == 403

    content = _build_workbook([VALID_ROW])
    resp = _upload(logged_in_client, content)
    assert resp.status_code == 403


def test_settings_hub_links_to_import_data(logged_in_client):
    resp = logged_in_client.get("/settings/")
    assert b'href="/import/"' in resp.data
