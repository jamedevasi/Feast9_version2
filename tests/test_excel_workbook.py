"""Capture template laid out like the earlier version's export ("Patients & Cases", one row
per case), and the Plan-B full Excel export, which must round-trip through the importer."""
import io
import json

from openpyxl import Workbook, load_workbook

from app import db
from app.excel_import import (
    CASES_SHEET_HEADERS, NO_CASES_SHEET_HEADERS, OLD_CASE_HEADERS, OLD_PATIENT_HEADERS,
)
from tests.conftest import get_csrf
from tests.test_roles import _create_user, _login, _logout

# The earlier version's export, exactly (sheet 'Patients & Cases').
OLD_EXPORT_HEADERS = [
    "Patient Name", "Age", "Sex", "Mobile", "Email", "Address", "Last Visited", "Emergency Contact",
    "EC Number", "Medical Conditions", "Allergies", "Case Title", "Status", "Doctor", "Procedures",
    "Total Cost (Rs.)", "Paid (Rs.)", "Balance (Rs.)", "Follow-up Date", "Next Action", "Case Opened",
    "Case Closed",
]


def _row(name="Anita Rao", mobile="9876543210", sex="Female", dpdp="Yes", headers=CASES_SHEET_HEADERS, **fields):
    row = dict.fromkeys(headers, "")
    row.update({"Patient Name": name, "Mobile": mobile, "Sex": sex, "DPDP Notice Accepted": dpdp})
    row.update({k.replace("_", " "): v for k, v in fields.items()})
    return [row[h] for h in headers]


def _case(title="RCT 36", doctor="Dr Test", **fields):
    return {"Case Title": title, "Doctor": doctor, **fields}


def _workbook(rows=(), no_case_rows=(), headers=CASES_SHEET_HEADERS):
    wb = Workbook()
    ws = wb.active
    ws.title = "Patients & Cases"
    ws.append(headers)
    for r in rows:
        ws.append(r)
    if no_case_rows:
        ws = wb.create_sheet("Patients (No Cases)")
        ws.append(NO_CASES_SHEET_HEADERS)
        for r in no_case_rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(client, content):
    token = get_csrf(client, "/backup/")
    return client.post(
        "/import/patients",
        data={"file": (io.BytesIO(content), "capture.xlsx"), "csrf_token": token},
        content_type="multipart/form-data",
    )


def _all(table):
    conn = db.get_db()
    rows = [dict(r) for r in conn.execute(f"SELECT * FROM {table} ORDER BY id")]
    conn.close()
    return rows


def _existing_patient(name="Existing Person", mobile="9876500000"):
    return db.add_patient({"name": name, "sex": "Male", "mobile": mobile, "dpdp_notice_accepted": 1})


def _existing_case(patient_id, doctor_id, title="Existing Case", cost=1000):
    return db.add_case({"patient_id": patient_id, "title": title, "doctor_id": doctor_id, "total_cost": cost})


def test_template_starts_with_the_old_export_columns_in_order(logged_in_client):
    db.add_doctor("Dr Test")
    wb = load_workbook(io.BytesIO(logged_in_client.get("/import/template.xlsx").data))
    ws = wb.active
    assert ws.title == "Patients & Cases"
    headers = [c.value for c in ws[1]]
    assert headers[:22] == OLD_EXPORT_HEADERS == OLD_PATIENT_HEADERS + OLD_CASE_HEADERS
    assert headers[22] == "DPDP Notice Accepted"  # the one worth filling for everyone, right after the pasted block
    assert headers[-2:] == ["Patient Ref", "Case Ref"]
    assert ws.max_row == 1  # no example rows that could be imported by accident

    fills = {c.value: c.fill.fgColor.rgb[-6:] for c in ws[1]}
    assert {fills["Patient Name"], fills["Sex"]} == {"F8D7DA"}  # red = required
    assert fills["DPDP Notice Accepted"] == "E2EFDA"  # recorded as entered; blank means not accepted
    assert {fills["Doctor"], fills["Guardian Name"], fills["Guardian Mobile"]} == {"FFF3CD"}  # amber = sometimes
    assert fills["Balance (Rs.)"] == "D9D9D9"  # grey = ignored
    assert fills["Email"] == "E2EFDA"

    assert "Dr Test" in [c.value for c in wb["Lists"]["A"]]
    assert "Read Me" in wb.sheetnames


def test_old_export_rows_import_with_only_dpdp_added(logged_in_client):
    """A row pasted from the earlier version's export — plus the DPDP column — imports: the
    same patient on two rows becomes one patient with two cases, Paid becomes a payment,
    'DR X' matches doctor 'Dr. X', and (Custom: ...) procedures survive their commas."""
    db.add_doctor("Dr. Jose Jimmy")
    db.add_procedure_type("RPD")
    rows = [
        _row("BEENA MICHAEL", "9847142169", **{"Age": 52, "Address": "ALUVA", "Last Visited": "2026-09-29",
             "Case Title": "RPD IRT 13,15,16", "Status": "Active", "Doctor": "DR JOSE JIMMY", "Procedures": "RPD",
             "Total Cost (Rs.)": 4000, "Paid (Rs.)": 1500, "Balance (Rs.)": 2500, "Case Opened": "2026-09-20"}),
        _row("Beena  Michael", "9847142169", **{"Case Title": "Filling", "Status": "Closed", "Doctor": "DR JOSE JIMMY",
             "Procedures": "Filling (GIC), (Custom: FILLING DONE IRT 16,26,36,46)", "Total Cost (Rs.)": 2400,
             # the earlier version wrote some Case Closed values with a time
             "Paid (Rs.)": 2400, "Case Opened": "2026-09-01", "Case Closed": "2026-09-10 07:10:27",
             "Last Visited": "2026-09-10"}),
    ]
    body = _upload(logged_in_client, _workbook(rows)).data.decode()
    assert "1</strong> patient(s) imported" in body
    assert "2</strong> case(s) imported" in body
    assert "2</strong> payment(s) imported" in body

    [patient] = _all("patients")
    assert patient["name"] == "BEENA MICHAEL" and patient["address"] == "ALUVA" and patient["age"] == 52
    assert patient["last_visited_date"] == "2026-09-29"  # the latest of the rows
    rpd, filling = _all("cases")
    assert json.loads(rpd["procedures_json"]) == ["RPD"]
    assert rpd["status"] == "Active" and rpd["created_at"].startswith("2026-09-20")
    assert filling["custom_procedure"] == "FILLING DONE IRT 16,26,36,46, Filling (GIC)"
    assert filling["status"] == "Closed" and filling["closed_at"].startswith("2026-09-10")
    assert db.get_case_balance(rpd["id"]) == 2500
    assert db.get_case_balance(filling["id"]) == 0
    payments = _all("payments")
    assert [p["payment_date"] for p in payments] == ["2026-09-20", "2026-09-10"]  # opened / closed date

    actions = [e["action"] for e in db.list_audit_log()]
    assert {"patients_bulk_imported", "cases_bulk_imported", "payment_added"} <= set(actions)


def test_old_export_without_dpdp_column_imports_everyone_as_not_accepted(logged_in_client):
    wb = Workbook()
    wb.active.title = "Patients & Cases"
    wb.active.append(OLD_EXPORT_HEADERS)
    row = {h: "" for h in OLD_EXPORT_HEADERS}
    row.update({"Patient Name": "Straight From Old Export", "Sex": "Male", "Mobile": "9000000099"})
    wb.active.append([row[h] for h in OLD_EXPORT_HEADERS])
    buf = io.BytesIO()
    wb.save(buf)
    _upload(logged_in_client, buf.getvalue())
    patient = db.list_patients()[0]
    assert patient["name"] == "Straight From Old Export" and patient["dpdp_notice_accepted"] == 0


def test_dpdp_no_never_gets_communications_consent(logged_in_client):
    rows = [_row("Said No", "9000000041", dpdp="No", **{"Communications Consent": "Yes"}),
            _row("Said Yes", "9000000042", dpdp="Yes", **{"Communications Consent": "Yes"})]
    _upload(logged_in_client, _workbook(rows))
    by_name = {p["name"]: p for p in db.list_patients()}
    assert (by_name["Said No"]["dpdp_notice_accepted"], by_name["Said No"]["comms_consent"]) == (0, 0)
    assert (by_name["Said Yes"]["dpdp_notice_accepted"], by_name["Said Yes"]["comms_consent"]) == (1, 1)
    assert by_name["Said Yes"]["dpdp_notice_accepted_at"] and not by_name["Said No"]["dpdp_notice_accepted_at"]


def test_patient_only_rows_and_no_cases_sheet(logged_in_client):
    no_case = _row("Walk In", "9123456780", headers=NO_CASES_SHEET_HEADERS)
    _upload(logged_in_client, _workbook([_row("Just Registered", "9123456781")], [no_case]))
    assert sorted(p["name"] for p in db.list_patients()) == ["Just Registered", "Walk In"]
    assert _all("cases") == []


def test_mandatory_fields_reported_and_dependent_rows_skipped(logged_in_client):
    db.add_doctor("Dr Test")
    rows = [
        _row("No Sex", "9000000001", sex="", **_case()),
        _row("No Sex", "9000000001", sex="", **_case("Second case")),
        _row("No Doctor", "9000000003", **_case(doctor="")),
        _row("Unknown Doctor", "9000000004", **_case(doctor="Dr Nobody")),
        _row("Minor", "9000000005", **{"Age": 12}),
        _row("Case Bits Without Title", "9000000006", **{"Doctor": "Dr Test", "Total Cost (Rs.)": 100}),
    ]
    body = _upload(logged_in_client, _workbook(rows)).data.decode()
    for message in ("Its patient was skipped (see Patients &amp; Cases row 2)",
                    "Sex must be Male or Female", "Doctor is required for a case",
                    "Doctor &#39;Dr Nobody&#39; isn&#39;t an active doctor", "Guardian name is required",
                    "Case Title is required when the row has case details"):
        assert message in body
    # Patients whose own details were fine are still imported, just without the bad case.
    assert sorted(p["name"] for p in db.list_patients()) == ["Case Bits Without Title", "No Doctor", "Unknown Doctor"]
    assert _all("cases") == []


def test_existing_patient_is_not_duplicated(logged_in_client):
    pid = _existing_patient("Existing Person", "9876500000")
    body = _upload(logged_in_client, _workbook([_row("existing person", "+91 98765 00000")])).data.decode()
    assert f"already in Feast9 as patient {pid}" in body
    assert len(db.list_patients()) == 1


def test_new_case_for_existing_patient_by_ref(logged_in_client):
    db.add_doctor("Dr Test")
    pid = _existing_patient("Existing Person")
    _upload(logged_in_client, _workbook([_row("Existing Person", **_case(), **{"Patient Ref": pid})]))
    assert len(db.list_patients()) == 1
    [case] = _all("cases")
    assert case["patient_id"] == pid


def test_ref_mistakes_are_caught(logged_in_client):
    doctor_id = db.add_doctor("Dr Test")
    pid = _existing_patient("Existing Person")
    other = _existing_patient("Other Person", "9876500001")
    case_id = _existing_case(other, doctor_id)
    body = _upload(logged_in_client, _workbook([
        _row("Somebody Else", **{"Patient Ref": pid}),
        _row("Existing Person", **{"Patient Ref": 999}),
        _row("Existing Person", **_case(), **{"Patient Ref": pid, "Case Ref": case_id}),
    ])).data.decode()
    assert f"Patient Ref {pid} is Existing Person, not Somebody Else" in body
    assert "Patient Ref 999 isn&#39;t a Feast9 patient ID" in body
    assert f"Case {case_id} belongs to patient {other}" in body
    assert len(_all("cases")) == 1


def test_invalid_values_reported(logged_in_client):
    db.add_doctor("Dr Test")
    body = _upload(logged_in_client, _workbook([_row(**_case(), **{
        "Total Cost (Rs.)": "abc", "Paid (Rs.)": -5, "Case Opened": "31/12/2026", "Status": "Pending",
    })])).data.decode()
    for message in ("Total Cost must be a number", "Paid must be a number", "isn&#39;t a valid date",
                    "Status must be Active or Closed"):
        assert message in body
    assert _all("cases") == []


def test_export_layout_and_audit(logged_in_client):
    doctor_id = db.add_doctor("Dr Test")
    pid = db.add_patient({"name": "Export Person", "sex": "Female", "mobile": "9876543210",
                          "allergies_json": '["Penicillin"]', "allergies_other": "Dust", "dpdp_notice_accepted": 1})
    lonely = _existing_patient("No Case Person", "9876500009")
    case_id = _existing_case(pid, doctor_id, "RCT 36", cost=5000)
    db.add_payment(case_id, pid, "2026-09-01", 1500, "Cash", "", "")
    db.add_visit_note(case_id, pid, "=HYPERLINK(\"http://evil\")", "2026-09-01")

    resp = logged_in_client.get("/import/export.xlsx")
    assert resp.status_code == 200 and "feast9_export_" in resp.headers["Content-Disposition"]
    wb = load_workbook(io.BytesIO(resp.data))
    for sheet in ("Read Me", "Upcoming Appointments", "Patients & Cases", "Patients (No Cases)", "Payments",
                  "All Appointments", "Visit Notes", "Prescriptions", "Lab Requisitions"):
        assert sheet in wb.sheetnames
    assert wb.active.title == "Upcoming Appointments"

    cases = list(wb["Patients & Cases"].iter_rows(values_only=True))
    assert list(cases[0][:22]) == OLD_EXPORT_HEADERS
    row = dict(zip(cases[0], cases[1]))
    assert (row["Patient Name"], row["Allergies"], row["Mobile"]) == ("Export Person", "Penicillin, Dust", "9876543210")
    assert (row["Paid (Rs.)"], row["Balance (Rs.)"], row["Doctor"]) == (1500, 3500, "Dr Test")
    assert (row["Patient Ref"], row["Case Ref"], row["DPDP Notice Accepted"]) == (pid, case_id, "Yes")

    no_cases = list(wb["Patients (No Cases)"].iter_rows(values_only=True))
    assert [dict(zip(no_cases[0], r))["Patient Ref"] for r in no_cases[1:]] == [lonely]

    note_cell = wb["Visit Notes"]["D2"]
    assert note_cell.data_type == "s" and note_cell.value.startswith("=HYPERLINK")  # never a live formula

    assert "full_data_exported" in [e["action"] for e in db.list_audit_log()]


def test_export_reimports_as_no_op_and_picks_up_plan_b_changes(logged_in_client):
    doctor_id = db.add_doctor("Dr Test")
    pid = _existing_patient("Existing Person")
    _existing_patient("No Case Person", "9876500009")
    case_id = _existing_case(pid, doctor_id)
    db.add_payment(case_id, pid, "2026-09-01", 200, "Cash", "", "")

    exported = logged_in_client.get("/import/export.xlsx").data
    body = _upload(logged_in_client, exported).data.decode()
    assert "Already in Feast9, left as they were: 2 patient(s), 1 case(s)" in body
    assert "No rows were skipped" in body
    assert len(db.list_patients()) == 2 and len(_all("cases")) == 1 and len(_all("payments")) == 1

    # Plan B: a walk-in with a case, and a payment on the existing case (Paid raised 200 -> 500).
    wb = load_workbook(io.BytesIO(exported))
    ws = wb["Patients & Cases"]
    headers = [c.value for c in ws[1]]
    ws.cell(row=2, column=headers.index("Paid (Rs.)") + 1, value=500)
    ws.cell(row=2, column=headers.index("Payment Method") + 1, value="UPI")
    ws.append(_row("Walk In", "9123456780", **_case("Scaling", **{"Paid (Rs.)": 300})))
    buf = io.BytesIO()
    wb.save(buf)
    _upload(logged_in_client, buf.getvalue())

    assert sorted(p["name"] for p in db.list_patients()) == ["Existing Person", "No Case Person", "Walk In"]
    assert len(_all("cases")) == 2
    assert db.get_case_balance(case_id) == 500
    top_up = [p for p in _all("payments") if p["case_id"] == case_id][-1]
    assert (top_up["amount"], top_up["method"]) == (300, "UPI")

    # Importing the same edited file again adds nothing more.
    _upload(logged_in_client, buf.getvalue())
    assert len(_all("payments")) == 3


def test_lowering_paid_on_an_exported_row_changes_nothing(logged_in_client):
    doctor_id = db.add_doctor("Dr Test")
    pid = _existing_patient("Existing Person")
    case_id = _existing_case(pid, doctor_id)
    db.add_payment(case_id, pid, "2026-09-01", 400, "Cash", "", "")
    wb = load_workbook(io.BytesIO(logged_in_client.get("/import/export.xlsx").data))
    ws = wb["Patients & Cases"]
    headers = [c.value for c in ws[1]]
    ws.cell(row=2, column=headers.index("Paid (Rs.)") + 1, value=100)
    buf = io.BytesIO()
    wb.save(buf)
    body = _upload(logged_in_client, buf.getvalue()).data.decode()
    assert "payments can&#39;t be reduced" in body
    assert db.get_case_balance(case_id) == 600


def test_export_requires_reauth(logged_in_client):
    with logged_in_client.session_transaction() as sess:
        sess["reauth_at"] = "2020-01-01T00:00:00"
    resp = logged_in_client.get("/import/export.xlsx")
    assert resp.status_code == 302 and "/reauth" in resp.headers["Location"]
    assert "full_data_exported" not in [e["action"] for e in db.list_audit_log()]


def test_export_is_admin_only(logged_in_client):
    _create_user(logged_in_client, "exportdoc", "doctor")
    _create_user(logged_in_client, "exportrecep", "receptionist")
    for username in ("exportdoc", "exportrecep"):
        _logout(logged_in_client)
        _login(logged_in_client, username)
        assert logged_in_client.get("/import/export.xlsx").status_code == 403
