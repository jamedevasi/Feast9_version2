"""Bulk patient import from an .xlsx spreadsheet (feast9_v2_agents.md §2/§3
excel_import.py: "Bulk 8,000-row xlsx import"). Every row is validated against the
exact same rules the manual patient-registration form enforces (name, sex, DPDP notice
acceptance, mobile format, guardian-for-minors) — bulk import never gets to skip DPDP
consent just because it's bulk. Valid rows are inserted in one transaction
(db.bulk_add_patients); invalid rows are skipped and reported back, never partially
written. This module never touches the database directly except through db.py.
"""
import datetime
import io

from openpyxl import Workbook, load_workbook

from app import db
from app.constants import MAX_IMPORT_ROWS, SEX_OPTIONS
from app.validators import compute_age, is_valid_mobile, normalize_date, today_iso

TEMPLATE_HEADERS = [
    "Name", "Date of Birth", "Age", "Sex", "Mobile", "Email", "Address",
    "Medical Conditions", "Allergies", "Is Pregnant", "Is Nursing",
    "Emergency Contact Name", "Emergency Contact Relation", "Emergency Contact Number",
    "DPDP Notice Accepted", "DPDP Notice Accepted Date", "Communications Consent",
    "Guardian Name", "Guardian Relation", "Guardian Mobile",
]
_REQUIRED_HEADERS = ("Name", "Sex", "DPDP Notice Accepted")

_TEMPLATE_EXAMPLE_ROW = [
    "Ananya Sharma", "1990-04-12", "", "Female", "9876543210", "ananya@example.com",
    "12 MG Road, Bengaluru", "Hypertension", "Penicillin", "No", "No",
    "Rohit Sharma", "Spouse", "9876500000",
    "Yes", "2020-01-15", "No", "", "", "",
]


def build_template_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "Patients"
    ws.append(TEMPLATE_HEADERS)
    ws.append(_TEMPLATE_EXAMPLE_ROW)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _yes(value):
    return value.strip().lower() in ("yes", "y", "true", "1")


def _cell_text(value):
    """openpyxl hands back native types per cell format — a date column comes back as
    datetime.date/datetime, a numeric column (mobile numbers are all-digit) as int/float —
    normalize everything to the plain string the rest of validation expects."""
    if value is None:
        return ""
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


class ImportResult:
    def __init__(self):
        self.imported_count = 0
        self.skipped = []  # [{"row": int, "name": str, "errors": [str, ...]}, ...]
        self.header_error = ""

    @property
    def ok(self):
        return not self.header_error


def _parse_row(row_values, header_index):
    def get(col):
        idx = header_index.get(col)
        if idx is None or idx >= len(row_values):
            return ""
        return _cell_text(row_values[idx])

    dob = normalize_date(get("Date of Birth"))
    age_computed = compute_age(dob) if dob else None
    manual_age_raw = get("Age")
    manual_age = int(manual_age_raw) if manual_age_raw.isdigit() else None
    age = age_computed if age_computed is not None else manual_age

    sex = get("Sex").capitalize()
    mobile = get("Mobile")
    dpdp_accepted = _yes(get("DPDP Notice Accepted"))
    dpdp_date = normalize_date(get("DPDP Notice Accepted Date")) or (today_iso() if dpdp_accepted else "")
    comms_consent = _yes(get("Communications Consent"))
    guardian_name = get("Guardian Name")
    guardian_mobile = get("Guardian Mobile")

    data = {
        "name": get("Name"),
        "date_of_birth": dob,
        "age": age,
        "sex": sex,
        "mobile": mobile,
        "email": get("Email"),
        "address": get("Address"),
        "medical_conditions_json": "[]",
        "medical_conditions_other": get("Medical Conditions"),
        "is_pregnant": _yes(get("Is Pregnant")),
        "is_nursing": _yes(get("Is Nursing")),
        "allergies_json": "[]",
        "allergies_other": get("Allergies"),
        "emergency_contact_name": get("Emergency Contact Name"),
        "emergency_contact_relation": get("Emergency Contact Relation"),
        "emergency_contact_number": get("Emergency Contact Number"),
        "dpdp_notice_accepted": dpdp_accepted,
        "dpdp_notice_accepted_at": dpdp_date,
        "comms_consent": comms_consent,
        "comms_consent_at": today_iso() if comms_consent else "",
        "guardian_name": guardian_name,
        "guardian_relation": get("Guardian Relation"),
        "guardian_mobile": guardian_mobile,
    }

    errors = []
    if not data["name"]:
        errors.append("Name is required.")
    if sex not in SEX_OPTIONS:
        errors.append("Sex must be Male or Female.")
    if not dpdp_accepted:
        errors.append("DPDP Notice Accepted must be Yes to import this patient — bulk import never bypasses consent.")
    if mobile and not is_valid_mobile(mobile):
        errors.append("Mobile number looks invalid — enter a 10-digit Indian mobile number.")
    if age is not None and age < 18:
        if not guardian_name:
            errors.append("Guardian name is required for patients under 18.")
        if not guardian_mobile or not is_valid_mobile(guardian_mobile):
            errors.append("A valid guardian mobile number is required for patients under 18.")

    return data, errors


def import_patients(file_bytes, actor=None):
    result = ImportResult()
    try:
        wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception:
        result.header_error = "Could not read this file — make sure it's a valid .xlsx spreadsheet."
        return result

    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        result.header_error = "The spreadsheet is empty."
        return result

    header_index = {}
    for i, cell in enumerate(header_row):
        if cell is not None:
            header_index[str(cell).strip()] = i

    missing = [h for h in _REQUIRED_HEADERS if h not in header_index]
    if missing:
        result.header_error = (
            f"Missing required column(s): {', '.join(missing)}. "
            "Download the template and use its exact column headers."
        )
        return result

    valid_rows = []
    row_number = 1
    for row_values in rows_iter:
        row_number += 1
        if row_number - 1 > MAX_IMPORT_ROWS:
            result.skipped.append({
                "row": row_number, "name": "",
                "errors": [f"Exceeds the {MAX_IMPORT_ROWS}-row limit — split the file and import the rest separately."],
            })
            break
        if row_values is None or all(v is None for v in row_values):
            continue  # blank row

        data, errors = _parse_row(row_values, header_index)
        if errors:
            result.skipped.append({"row": row_number, "name": data["name"], "errors": errors})
        else:
            valid_rows.append(data)

    if valid_rows:
        db.bulk_add_patients(valid_rows, actor=actor)
        result.imported_count = len(valid_rows)

    return result
