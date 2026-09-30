"""Plan-B Excel export — the whole practice in one .xlsx, for when Feast9 itself is
unavailable. Its "Patients & Cases" and "Patients (No Cases)" sheets keep the earlier
version's export layout (app/excel_import.py) with each row's Feast9 ids in Patient Ref /
Case Ref: the clinic keeps working in this file during an outage — new rows for new
patients and cases, a higher Paid figure for a payment taken — and imports it once the
system is back. Exported rows are recognised by id, so only what changed gets written.
The other sheets (appointments, payments, visit notes, prescriptions, lab requisitions)
are read-only reference.
"""
import datetime
import io
import json

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from app import db
from app.constants import DEFAULT_CLINIC_NAME
from app.excel_import import (
    CASES_SHEET, CASES_SHEET_HEADERS, NO_CASES_SHEET, NO_CASES_SHEET_HEADERS,
    add_dropdowns, set_active_sheet, style_header_row, write_read_me,
)
from app.validators import compute_age, today_iso

# Same row colours as the earlier version's export.
_ACTIVE_FILL = PatternFill("solid", fgColor="E2EFDA")
_DUE_TODAY_FILL = PatternFill("solid", fgColor="FFF3CD")
_OVERDUE_FILL = PatternFill("solid", fgColor="F8D7DA")

_APPOINTMENT_HEADERS = [
    "Date", "Start", "End", "Patient Name", "Mobile", "Doctor", "Title", "Status", "Recurring", "Notes",
    "Patient Ref", "Case Ref",
]


def _yn(value):
    return "Yes" if value else "No"


def _date(value):
    return (value or "")[:10]


def _join_list(json_text, other):
    try:
        items = json.loads(json_text or "[]")
    except ValueError:
        items = []
    return ", ".join([*items, *([other] if other else [])])


def _append(ws, values):
    """Appends a row with every text value kept as literal text: openpyxl would otherwise
    write a string starting with '=' as a live formula — patient-typed notes must never
    become executable in Excel."""
    ws.append([None if v == "" else v for v in values])
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def _fill_row(ws, fill):
    for cell in ws[ws.max_row]:
        cell.fill = fill


def _sheet(wb, title, headers, capture=False):
    ws = wb.create_sheet(title)
    ws.append(headers)
    style_header_row(ws, headers, capture=capture)
    return ws


def _finish(ws):
    ws.auto_filter.ref = ws.dimensions


def _patient_values(p):
    """Old-layout patient columns, then this version's extra patient columns."""
    age = compute_age(p.get("date_of_birth")) if p.get("date_of_birth") else p.get("age")
    old = [
        p["name"], age if age is not None else "", p.get("sex") or "", p.get("mobile") or "",
        p.get("email") or "", p.get("address") or "",
        p.get("last_visit_note_date") or p.get("last_visited_date") or "",
        p.get("emergency_contact_name") or "", p.get("emergency_contact_number") or "",
        _join_list(p.get("medical_conditions_json"), p.get("medical_conditions_other")),
        _join_list(p.get("allergies_json"), p.get("allergies_other")),
    ]
    extra = [
        _yn(p.get("dpdp_notice_accepted")), _date(p.get("dpdp_notice_accepted_at")),
        p.get("date_of_birth") or "", _yn(p.get("comms_consent")),
        _yn(p.get("is_pregnant")), _yn(p.get("is_nursing")), p.get("emergency_contact_relation") or "",
        p.get("guardian_name") or "", p.get("guardian_relation") or "", p.get("guardian_mobile") or "",
    ]
    return old, extra


def _appointment_row(a):
    return [
        a["appt_date"], a["start_time"], a.get("end_time") or "", a["patient_name"], a.get("patient_mobile") or "",
        a.get("doctor_name") or "", a.get("title") or "", a["status"], _yn(a.get("series_id")),
        a.get("notes") or "", a["patient_id"], a.get("case_id") or "",
    ]


def build_export_xlsx(generated_by=""):
    snap = db.get_export_snapshot()
    today = today_iso()
    wb = Workbook()
    read_me = wb.active
    read_me.title = "Read Me"
    patients = {p["id"]: p for p in snap["patients"]}

    # Upcoming appointments first — what a front desk needs on the first morning of an outage.
    upcoming = [a for a in snap["appointments"] if a["appt_date"] >= today and a["status"] == "Scheduled"]
    ws = _sheet(wb, "Upcoming Appointments", _APPOINTMENT_HEADERS)
    for a in upcoming:
        _append(ws, _appointment_row(a))
    _finish(ws)

    ws = _sheet(wb, CASES_SHEET, CASES_SHEET_HEADERS, capture=True)
    for c in snap["cases"]:
        old_patient, extra_patient = _patient_values(patients[c["patient_id"]])
        try:
            procedures = ", ".join(json.loads(c.get("procedures_json") or "[]"))
        except ValueError:
            procedures = ""
        total = c.get("total_cost") or 0
        _append(ws, old_patient + [
            c["title"], c["status"], c.get("doctor_name") or "", procedures, total, c["paid"], total - c["paid"],
            c.get("follow_up_date") or "", c.get("next_action_note") or "",
            _date(c.get("created_at")), _date(c.get("closed_at")),
        ] + extra_patient + [c.get("custom_procedure") or "", "", "", c["patient_id"], c["id"]])
        if c["status"] == "Active":
            follow_up = c.get("follow_up_date") or ""
            _fill_row(ws, _OVERDUE_FILL if follow_up and follow_up < today
                      else _DUE_TODAY_FILL if follow_up == today else _ACTIVE_FILL)
    _finish(ws)

    with_cases = {c["patient_id"] for c in snap["cases"]}
    ws = _sheet(wb, NO_CASES_SHEET, NO_CASES_SHEET_HEADERS, capture=True)
    for p in snap["patients"]:
        if p["id"] not in with_cases:
            old_patient, extra_patient = _patient_values(p)
            _append(ws, old_patient + extra_patient + [p["id"]])
    _finish(ws)

    ws = _sheet(wb, "Payments", ["Payment Date", "Patient Name", "Case Title", "Amount (Rs.)", "Method",
                                 "Reference", "Notes", "Patient Ref", "Case Ref"])
    for pay in snap["payments"]:
        _append(ws, [pay["payment_date"], pay["patient_name"], pay["case_title"], pay["amount"],
                     pay.get("method") or "", pay.get("reference") or "", pay.get("notes") or "",
                     pay["patient_id"], pay["case_id"]])
    _finish(ws)

    ws = _sheet(wb, "All Appointments", _APPOINTMENT_HEADERS)
    for a in snap["appointments"]:
        _append(ws, _appointment_row(a))
    _finish(ws)

    ws = _sheet(wb, "Visit Notes", ["Visit Date", "Patient Name", "Case Title", "Note", "Patient Ref", "Case Ref"])
    for v in snap["visit_notes"]:
        _append(ws, [v["visit_date"], v["patient_name"], v["case_title"], v["note"], v["patient_id"], v["case_id"]])
    _finish(ws)

    ws = _sheet(wb, "Prescriptions", ["Date", "Patient Name", "Case Title", "Prescription", "Patient Ref", "Case Ref"])
    for rx in snap["prescriptions"]:
        _append(ws, [rx["prescribed_date"], rx["patient_name"], rx["case_title"], rx["rx_details"],
                     rx["patient_id"], rx["case_id"]])
    _finish(ws)

    ws = _sheet(wb, "Lab Requisitions", [
        "Sent", "Patient Name", "Case Title", "Lab", "Work", "Expected Return", "Received", "Status", "Notes",
        "Patient Ref", "Case Ref",
    ])
    for lab in snap["lab_requisitions"]:
        _append(ws, [
            lab["sent_date"], lab["patient_name"], lab["case_title"], lab.get("lab_name") or "",
            lab["work_description"], lab.get("expected_return") or "", lab.get("received_date") or "",
            lab["status"], lab.get("notes") or "", lab["patient_id"], lab["case_id"],
        ])
    _finish(ws)

    clinic = db.get_setting("clinic_name") or DEFAULT_CLINIC_NAME
    generated = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    write_read_me(read_me, intro_rows=[
        (f"{clinic} — full data export", None),
        ("Generated", f"{generated}{f' by {generated_by}' if generated_by else ''}"),
        ("Contains", f"{len(snap['patients'])} patients, {len(snap['cases'])} cases, {len(snap['payments'])} payments, "
                     f"{len(snap['appointments'])} appointments ({len(upcoming)} upcoming), "
                     f"{len(snap['visit_notes'])} visit notes, {len(snap['prescriptions'])} prescriptions, "
                     f"{len(snap['lab_requisitions'])} lab requisitions."),
        ("Confidential", "Full patient health and financial records. Keep it on a secured device, don't email or "
                         "share it, and delete it once it's no longer needed (DPDP Act 2023)."),
        ("", ""),
        ("If Feast9 is down", None),
        ("Look up", "Upcoming Appointments, Patients & Cases (allergies and medical conditions are in the patient "
                    "columns), Payments, Visit Notes, Prescriptions."),
        ("New patient or case", "Add a row at the bottom of 'Patients & Cases' and leave Patient Ref and Case Ref "
                                "blank. For a new case for an existing patient, copy their row, keep the Patient "
                                "Ref, clear the case columns and Case Ref, and fill in the new case."),
        ("Payment taken", "Raise the case's Paid figure by the amount received. Put the date in Payment Date and "
                          "the method in Payment Method."),
        ("When Feast9 is back", "Import this file from Settings > Backup & Data. Only what you added or raised is "
                                "recorded. Appointments, visit notes and prescriptions from the outage have to be "
                                "entered in Feast9 by hand."),
        ("Row colours", "Green = active case. Amber = follow-up due today. Red = follow-up overdue."),
        ("", ""),
    ])
    add_dropdowns(wb, [(CASES_SHEET, CASES_SHEET_HEADERS), (NO_CASES_SHEET, NO_CASES_SHEET_HEADERS)])
    set_active_sheet(wb, "Upcoming Appointments")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), snap
